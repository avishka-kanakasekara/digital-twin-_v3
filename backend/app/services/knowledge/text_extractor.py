"""
Text Extractor — Extracts text from supported document formats.

Supported formats:
    PDF  — extracts text per page, detects image-only PDFs
    DOCX — extracts paragraphs, headings, tables
    TXT/MD — direct read
    JSON — safe parse
    CSV  — safe parse

Never silently produces empty output.
"""
from __future__ import annotations

import csv
import io
import json
import os
from dataclasses import dataclass, field
from pathlib import Path


# ── Result Types ──────────────────────────────────────────────

@dataclass
class ExtractedPage:
    page_number: int
    text: str
    char_count: int = 0

    def __post_init__(self):
        self.char_count = len(self.text)


@dataclass
class ExtractedDocument:
    full_text: str
    pages: list[ExtractedPage] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)
    is_image_only: bool = False
    extraction_method: str = "unknown"
    word_count: int = 0
    char_count: int = 0

    def __post_init__(self):
        self.word_count = len(self.full_text.split())
        self.char_count = len(self.full_text)


@dataclass
class ExtractionResult:
    success: bool
    document: ExtractedDocument | None = None
    error_code: str | None = None
    error_message: str | None = None


# ── Main Extractor ────────────────────────────────────────────

def extract_text(file_path: str, extension: str) -> ExtractionResult:
    """
    Extract text from a stored file.
    
    Args:
        file_path: Absolute path to the stored file
        extension: Normalized extension (.pdf, .docx, .txt, .md, .json, .csv)
    
    Returns:
        ExtractionResult with populated ExtractedDocument or error details
    """
    if not Path(file_path).exists():
        return ExtractionResult(
            success=False,
            error_code="FILE_NOT_FOUND",
            error_message="Stored file could not be located for extraction.",
        )

    try:
        if extension == ".pdf":
            return _extract_pdf(file_path)
        elif extension == ".docx":
            return _extract_docx(file_path)
        elif extension in {".txt", ".md"}:
            return _extract_text_file(file_path)
        elif extension == ".json":
            return _extract_json(file_path)
        elif extension == ".csv":
            return _extract_csv(file_path)
        else:
            return ExtractionResult(
                success=False,
                error_code="UNSUPPORTED_FORMAT",
                error_message=f"Extraction not implemented for '{extension}'.",
            )
    except Exception as exc:
        return ExtractionResult(
            success=False,
            error_code="TEXT_EXTRACTION_FAILED",
            error_message=f"Extraction failed: {str(exc)[:200]}",
        )


def extract_text_from_bytes(content: bytes, extension: str) -> ExtractionResult:
    """Extract text directly from bytes without requiring a stored file."""
    try:
        if extension == ".pdf":
            return _extract_pdf_bytes(content)
        elif extension == ".docx":
            return _extract_docx_bytes(content)
        elif extension in {".txt", ".md"}:
            text = content.decode("utf-8", errors="replace")
            return _build_plaintext_result(text, "plain_text")
        elif extension == ".json":
            parsed = json.loads(content.decode("utf-8", errors="replace"))
            text = json.dumps(parsed, indent=2)
            return _build_plaintext_result(text, "json")
        elif extension == ".csv":
            text = _csv_bytes_to_text(content)
            return _build_plaintext_result(text, "csv")
        else:
            return ExtractionResult(
                success=False,
                error_code="UNSUPPORTED_FORMAT",
                error_message=f"Extraction not implemented for '{extension}'.",
            )
    except Exception as exc:
        return ExtractionResult(
            success=False,
            error_code="TEXT_EXTRACTION_FAILED",
            error_message=f"Extraction failed: {str(exc)[:200]}",
        )


# ── PDF Extraction ────────────────────────────────────────────

def _extract_pdf(file_path: str) -> ExtractionResult:
    with open(file_path, "rb") as f:
        content = f.read()
    return _extract_pdf_bytes(content)


def _extract_pdf_bytes(content: bytes) -> ExtractionResult:
    try:
        import PyPDF2
    except ImportError:
        return ExtractionResult(
            success=False,
            error_code="DEPENDENCY_MISSING",
            error_message="PyPDF2 is not installed. Run: pip install PyPDF2",
        )

    try:
        reader = PyPDF2.PdfReader(io.BytesIO(content))
        pages: list[ExtractedPage] = []
        all_text_parts: list[str] = []

        for page_num, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            page_text = page_text.strip()
            pages.append(ExtractedPage(page_number=page_num, text=page_text))
            if page_text:
                all_text_parts.append(page_text)

        full_text = "\n\n".join(all_text_parts).strip()
        is_image_only = len(full_text) < 50 and len(reader.pages) > 0

        if is_image_only:
            return ExtractionResult(
                success=False,
                error_code="OCR_REQUIRED",
                error_message="This PDF appears to be image-only or scanned. Text extraction requires OCR, which is not currently available.",
            )

        if not full_text:
            return ExtractionResult(
                success=False,
                error_code="DOCUMENT_UNREADABLE",
                error_message="No readable text could be extracted from this PDF.",
            )

        return ExtractionResult(
            success=True,
            document=ExtractedDocument(
                full_text=full_text,
                pages=pages,
                metadata={
                    "page_count": len(reader.pages),
                    "info": {k: str(v) for k, v in (reader.metadata or {}).items()},
                },
                extraction_method="PyPDF2",
            ),
        )
    except PyPDF2.errors.PdfReadError as exc:
        return ExtractionResult(
            success=False,
            error_code="DOCUMENT_UNREADABLE",
            error_message=f"PDF could not be read: {str(exc)[:200]}",
        )


# ── DOCX Extraction ───────────────────────────────────────────

def _extract_docx(file_path: str) -> ExtractionResult:
    with open(file_path, "rb") as f:
        content = f.read()
    return _extract_docx_bytes(content)


def _extract_docx_bytes(content: bytes) -> ExtractionResult:
    try:
        import docx as python_docx
    except ImportError:
        return ExtractionResult(
            success=False,
            error_code="DEPENDENCY_MISSING",
            error_message="python-docx is not installed. Run: pip install python-docx",
        )

    try:
        doc = python_docx.Document(io.BytesIO(content))
        parts: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)

        # Extract table content
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)

        full_text = "\n".join(parts).strip()

        if not full_text:
            return ExtractionResult(
                success=False,
                error_code="DOCUMENT_UNREADABLE",
                error_message="No readable text could be extracted from this DOCX file.",
            )

        # Simulate page structure (DOCX doesn't have explicit pages)
        pages = [ExtractedPage(page_number=1, text=full_text)]

        return ExtractionResult(
            success=True,
            document=ExtractedDocument(
                full_text=full_text,
                pages=pages,
                metadata={"paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)},
                extraction_method="python-docx",
            ),
        )
    except Exception as exc:
        return ExtractionResult(
            success=False,
            error_code="TEXT_EXTRACTION_FAILED",
            error_message=f"DOCX extraction failed: {str(exc)[:200]}",
        )


# ── Plain Text / Markdown ─────────────────────────────────────

def _extract_text_file(file_path: str) -> ExtractionResult:
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    return _build_plaintext_result(text, "plain_text")


# ── JSON ──────────────────────────────────────────────────────

def _extract_json(file_path: str) -> ExtractionResult:
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        raw = f.read()
    try:
        parsed = json.loads(raw)
        text = json.dumps(parsed, indent=2)
    except json.JSONDecodeError as exc:
        return ExtractionResult(
            success=False,
            error_code="DOCUMENT_UNREADABLE",
            error_message=f"JSON file is malformed: {str(exc)[:200]}",
        )
    return _build_plaintext_result(text, "json")


# ── CSV ───────────────────────────────────────────────────────

def _extract_csv(file_path: str) -> ExtractionResult:
    with open(file_path, "rb") as f:
        content = f.read()
    text = _csv_bytes_to_text(content)
    return _build_plaintext_result(text, "csv")


def _csv_bytes_to_text(content: bytes) -> str:
    decoded = content.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(decoded))
    rows = [" | ".join(row) for row in reader if any(cell.strip() for cell in row)]
    return "\n".join(rows)


# ── Helpers ───────────────────────────────────────────────────

def _build_plaintext_result(text: str, method: str) -> ExtractionResult:
    text = text.strip()
    if not text:
        return ExtractionResult(
            success=False,
            error_code="DOCUMENT_UNREADABLE",
            error_message="The document contains no readable text.",
        )
    pages = [ExtractedPage(page_number=1, text=text)]
    return ExtractionResult(
        success=True,
        document=ExtractedDocument(
            full_text=text,
            pages=pages,
            metadata={},
            extraction_method=method,
        ),
    )
