"""
Extract Module 3 (Scenario Simulator) and related tables from Organization_Twin_Dummy_Data_M1-M9.xlsx
and save as backend/data/scenario_data.json
"""

import json
from pathlib import Path
import openpyxl

EXCEL_PATH = Path("Organization_Twin_Dummy_Data_M1-M9.xlsx")
OUTPUT_PATH = Path("backend/data/scenario_data.json")

SHEETS = [
    "baseline_snapshot",
    "scenario",
    "scenario_variable",
    "promotion_audit_trail",
    "org_unit",
    "cost_band",
    "future_roles",
    "skills",
    "operational_causal_metrics"
]

def extract():
    print(f"Loading {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    
    extracted_data = {}
    
    for sheet_name in SHEETS:
        if sheet_name not in wb.sheetnames:
            print(f"Warning: Sheet {sheet_name} not found!")
            continue
            
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            extracted_data[sheet_name] = []
            continue
            
        headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
        
        sheet_rows = []
        for r in rows[1:]:
            if not any(r):
                continue
            row_dict = {}
            for i, val in enumerate(r):
                if i < len(headers):
                    header = headers[i]
                    if hasattr(val, "isoformat"):
                        val = val.isoformat()
                    row_dict[header] = val
            sheet_rows.append(row_dict)
            
        print(f"Extracted {len(sheet_rows)} rows from {sheet_name}")
        extracted_data[sheet_name] = sheet_rows
        
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(extracted_data, f, indent=2, ensure_ascii=False)
        
    print(f"Successfully saved to {OUTPUT_PATH}")

if __name__ == "__main__":
    extract()
