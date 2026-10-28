import openpyxl

wb = openpyxl.load_workbook('Organization_Twin_Dummy_Data_M1-M9.xlsx', data_only=True)
scenarios = list(wb['scenario'].iter_rows(values_only=True))
hdr = scenarios[0]
scen_list = [dict(zip(hdr, r)) for r in scenarios[1:] if any(r)]

vars_rows = list(wb['scenario_variable'].iter_rows(values_only=True))
v_hdr = vars_rows[0]
var_list = [dict(zip(v_hdr, r)) for r in vars_rows[1:] if any(r)]

for s in scen_list:
    s_vars = [v for v in var_list if v['scenario_id'] == s['scenario_id']]
    name = s['scenario_name']
    stype = s['scenario_type']
    ou = s['org_unit_id']
    status = s['promotion_status']
    print(f"{s['scenario_id']}: {name} [{stype}] ({ou}) Status: {status} Horizon: {s['horizon_start_fy']}-{s['horizon_end_fy']}")
    for v in s_vars:
        print(f"   * {v['variable_type']} -> target: {v['target_id']}, val: {v['value']}, ramp: {v['ramp_profile']}, cost: {v['unit_cost_lkr']}")

print("\n--- BASELINE SNAPSHOTS ---")
snaps = list(wb['baseline_snapshot'].iter_rows(values_only=True))
s_hdr = snaps[0]
for r in snaps[1:]:
    d = dict(zip(s_hdr, r))
    cost = d['run_cost_annual_lkr']
    print(f"{d['snapshot_id']}: OU={d['org_unit_id']} HC={d['headcount_current']} Budgeted={d['headcount_budgeted']} RunCost={cost} FTE_Ratio={d['fte_contractor_ratio']} Gap={d['capability_gap_at_snapshot']} Frozen={d['is_frozen']}")
